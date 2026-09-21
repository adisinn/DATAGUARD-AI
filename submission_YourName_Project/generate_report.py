from docx import Document
from pathlib import Path
import pandas as pd

DATA_DIR = Path('sample_data')
OUT = Path(__file__).parent

def create_report():
    df = pd.read_csv(DATA_DIR / 'customer_profiles_enriched.csv')
    doc = Document()
    doc.add_heading('YourName Project Report', level=1)
    doc.add_paragraph('Project: DataGuard AI — sample analysis and profiling demo')
    doc.add_heading('Dataset', level=2)
    doc.add_paragraph(str(DATA_DIR / 'customer_profiles_enriched.csv'))
    doc.add_heading('Summary Statistics', level=2)
    summary = df.describe(include='all').T
    for col in summary.index:
        doc.add_heading(col, level=3)
        stats = summary.loc[col].to_dict()
        for k, v in stats.items():
            doc.add_paragraph(f"{k}: {v}")
    out_path = OUT / 'YourName_ProjectReport.docx'
    doc.save(out_path)
    print('Saved', out_path)

if __name__ == '__main__':
    create_report()
